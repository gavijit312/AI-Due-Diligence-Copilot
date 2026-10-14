from pathlib import Path
from typing import List, Dict, Any
from pypdf import PdfReader
import docx
import openpyxl
from ingestion.ocr import ocr_pdf_fallback


def parse_pdf(file_path: str) -> List[Dict[str, Any]]:
    reader = PdfReader(file_path)
    pages = []
    needs_ocr = True

    for idx, page in enumerate(reader.pages, start=1):
        content = page.extract_text() or ""
        if len(content.strip()) > 50:
            needs_ocr = False
        pages.append({"page": idx, "text": content.strip()})

    if needs_ocr:
        return ocr_pdf_fallback(file_path)
    return pages


def parse_docx(file_path: str) -> List[Dict[str, Any]]:
    doc = docx.Document(file_path)
    full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    return [{"page": 1, "text": full_text}]


def parse_xlsx(file_path: str) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets_data = []

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows_str = []
        for row in ws.iter_rows(values_only=True):
            cleaned_row = [str(cell) for cell in row if cell is not None]
            if cleaned_row:
                rows_str.append(" | ".join(cleaned_row))
        sheets_data.append({"page": sheet, "text": "\n".join(rows_str)})

    return sheets_data


def parse_file(file_path: str) -> List[Dict[str, Any]]:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    elif ext in [".xlsx", ".xls"]:
        return parse_xlsx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")